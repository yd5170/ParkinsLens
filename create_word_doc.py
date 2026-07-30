import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def create_document():
    doc = Document()
    
    # A4 Page Setup with 2cm margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)
        section.page_width = Inches(8.27)  # A4 width
        section.page_height = Inches(11.69) # A4 height

    # Base Styles Setup
    normal_style = doc.styles['Normal']
    normal_style.font.name = '맑은 고딕'
    normal_style.font.size = Pt(10)
    normal_style.font.color.rgb = RGBColor(0x2D, 0x37, 0x48) # #2D3748

    # Helper Functions for Table Cell Styling
    def set_cell_background(cell, fill_color):
        tcPr = cell._element.get_or_add_tcPr()
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_color}"/>')
        tcPr.append(shd)

    def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
        tcPr = cell._element.get_or_add_tcPr()
        tcMar = parse_xml(f'''
            <w:tcMar {nsdecls("w")}>
                <w:top w:w="{top}" w:type="dxa"/>
                <w:bottom w:w="{bottom}" w:type="dxa"/>
                <w:left w:w="{left}" w:type="dxa"/>
                <w:right w:w="{right}" w:type="dxa"/>
            </w:tcMar>
        ''')
        tcPr.append(tcMar)

    def set_table_borders(table, color="D2D6DC"):
        tblPr = table._element.xpath('w:tblPr')
        if tblPr:
            borders = parse_xml(f'''
                <w:tblBorders {nsdecls("w")}>
                    <w:top w:val="single" w:sz="4" w:space="0" w:color="{color}"/>
                    <w:bottom w:val="single" w:sz="8" w:space="0" w:color="1A365D"/>
                    <w:insideH w:val="single" w:sz="4" w:space="0" w:color="{color}"/>
                    <w:insideV w:val="none"/>
                    <w:left w:val="none"/>
                    <w:right w:val="none"/>
                </w:tblBorders>
            ''')
            tblPr[0].append(borders)

    # 1. Main Title Header
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(4)
    run_title = title_p.add_run("MRI 기반 파킨슨병 조기 진단 AI 솔루션")
    run_title.font.name = '맑은 고딕'
    run_title.font.size = Pt(22)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D) # Deep Navy

    subtitle_p = doc.add_paragraph()
    subtitle_p.paragraph_format.space_after = Pt(16)
    run_sub = subtitle_p.add_run("Data Flow Diagram (DFD) 및 데이터 흐름 명세서")
    run_sub.font.name = '맑은 고딕'
    run_sub.font.size = Pt(14)
    run_sub.font.bold = True
    run_sub.font.color.rgb = RGBColor(0x4A, 0x55, 0x68) # Slate Gray

    # Horizontal Divider Line
    p_line = doc.add_paragraph()
    p_line.paragraph_format.space_after = Pt(16)
    p_line_border = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="12" w:space="1" w:color="2B6CB0"/></w:pBdr>')
    p_line._element.get_or_add_pPr().append(p_line_border)

    # 2. Section: 개요 (Overview)
    h1 = doc.add_paragraph()
    h1.paragraph_format.space_before = Pt(12)
    h1.paragraph_format.space_after = Pt(6)
    r_h1 = h1.add_run("1. 시스템 개요")
    r_h1.font.name = '맑은 고딕'
    r_h1.font.size = Pt(14)
    r_h1.font.bold = True
    r_h1.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)

    p_desc = doc.add_paragraph()
    p_desc.paragraph_format.space_after = Pt(12)
    p_desc.paragraph_format.line_spacing = 1.2
    p_desc.add_run(
        "본 문서는 뇌 MRI 데이터를 활용한 파킨슨병 조기 진단 AI 솔루션의 데이터 흐름(Data Flow)을 정의한 명세서입니다. "
        "2D-MRI 영상 데이터의 입력부터 5단계 자동 전처리, 3D 중첩학습 기반 AI 분석, 의심 병변 시각화(M3d-CAM), "
        "그리고 최종 판독 리포트 자동 생성 및 전달까지의 전 과정 데이터 처리 구조를 포함합니다."
    )

    # 3. Section: DFD 다이어그램 구조 (Diagram Architecture)
    h2 = doc.add_paragraph()
    h2.paragraph_format.space_before = Pt(14)
    h2.paragraph_format.space_after = Pt(6)
    r_h2 = h2.add_run("2. 데이터 플로우 다이어그램 (DFD Process Diagram)")
    r_h2.font.name = '맑은 고딕'
    r_h2.font.size = Pt(14)
    r_h2.font.bold = True
    r_h2.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)

    # Box container for DFD Text Flow
    box_table = doc.add_table(rows=1, cols=1)
    box_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = box_table.cell(0, 0)
    set_cell_background(cell, "F7FAFC") # Light cool gray
    set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
    
    # Border for DFD box
    tcPr = cell._element.get_or_add_tcPr()
    tcBorders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="6" w:space="0" w:color="CBD5E0"/>
            <w:left w:val="single" w:sz="18" w:space="0" w:color="2B6CB0"/>
            <w:bottom w:val="single" w:sz="6" w:space="0" w:color="CBD5E0"/>
            <w:right w:val="single" w:sz="6" w:space="0" w:color="CBD5E0"/>
        </w:tcBorders>
    ''')
    tcPr.append(tcBorders)

    dfd_text_lines = [
        ("의료진 / 사용자 (2D-MRI 데이터 입력)", True, "1A365D"),
        ("  │", False, "718096"),
        ("  ▼  [1. 2D DICOM 영상 업로드]", False, "2B6CB0"),
        ("STEP 1. 데이터 전처리 (5단계 자동화)", True, "2B6CB0"),
        ("  ├─ 1.1 DICOM → NIfTI 변환  │  1.2 두개골 제거", False, "4A5568"),
        ("  ├─ 1.3 뇌 영역 정규화     │  1.4 Intensity 정규화", False, "4A5568"),
        ("  └─ 1.5 3D 리사이즈 (56×56×56)", False, "4A5568"),
        ("  │", False, "718096"),
        ("  ▼  [2. 전처리 완료 3D MRI 데이터]", False, "2B6CB0"),
        ("STEP 2. AI 모델 분석 (심층 3D 중첩학습)", True, "2B6CB0"),
        ("  ├─ 2.1 3D-CNN (공간 특징 추출)   │  2.2 3D-ResNet (심층 특징 학습)", False, "4A5568"),
        ("  └─ 2.3 CCA (특징 결합)          │  2.4 WOA (모델 최적화)", False, "4A5568"),
        ("  │", False, "718096"),
        ("  ▼  [3. 파킨슨병 여부 예측 (정상 / 전구기 / 파킨슨병)]", False, "2B6CB0"),
        ("  ├─────────────── Condition: 판정 분기 ───────────────┤", False, "E53E3E"),
        ("  │  [가능성 높음] ➔ 고위험군 알림 발송                │", False, "C53030"),
        ("  │  [정상/가능성 낮음 or 재검토] ➔ 피드백 반영 / STEP 3 진행 │", False, "2F855A"),
        ("  └────────────────────────────────────────────────────┘", False, "E53E3E"),
        ("  │", False, "718096"),
        ("  ▼  [4. 의심 병변 및 분석 지표 데이터]", False, "2B6CB0"),
        ("STEP 3. 진단 결과 시각화 (M3d-CAM)", True, "2B6CB0"),
        ("  └─ 의심 병변 영역 시각화 및 히트맵 생성", False, "4A5568"),
        ("  │", False, "718096"),
        ("  ▼  [5. 시각화 + 요약 데이터]", False, "2B6CB0"),
        ("STEP 4. 판독 리포트 자동 생성", True, "2B6CB0"),
        ("  └─ 진단 결과 요약, 병변 시각화, 분석 지표 종합", False, "4A5568"),
        ("  │", False, "718096"),
        ("  ▼  [6. 최종 판독 리포트]", False, "2B6CB0"),
        ("STEP 5. 판독 리포트 제공", True, "1A365D"),
        ("  ├─ 의료진 다운로드 및 저장", False, "4A5568"),
        ("  └─ 리포트 DB 저장소 보관", False, "4A5568")
    ]

    p_cell = cell.paragraphs[0]
    p_cell.paragraph_format.space_before = Pt(0)
    p_cell.paragraph_format.space_after = Pt(0)
    p_cell.paragraph_format.line_spacing = 1.15

    for idx, (line, is_bold, color_hex) in enumerate(dfd_text_lines):
        if idx > 0:
            p_cell = cell.add_paragraph()
            p_cell.paragraph_format.space_before = Pt(0)
            p_cell.paragraph_format.space_after = Pt(0)
            p_cell.paragraph_format.line_spacing = 1.15

        run = p_cell.add_run(line)
        run.font.name = 'Consolas' if ('│' in line or '─' in line or '▼' in line or '├' in line or '└' in line) else '맑은 고딕'
        run.font.size = Pt(9.5)
        run.font.bold = is_bold
        r, g, b = int(color_hex[:2], 16), int(color_hex[2:4], 16), int(color_hex[4:], 16)
        run.font.color.rgb = RGBColor(r, g, b)

    # Spacer
    p_space = doc.add_paragraph()
    p_space.paragraph_format.space_after = Pt(12)

    # 4. Section: 데이터 흐름 명세 표 (Data Flow Specification Table)
    h3 = doc.add_paragraph()
    h3.paragraph_format.space_before = Pt(14)
    h3.paragraph_format.space_after = Pt(8)
    r_h3 = h3.add_run("3. 데이터 흐름 상세 명세표 (Data Flow Table)")
    r_h3.font.name = '맑은 고딕'
    r_h3.font.size = Pt(14)
    r_h3.font.bold = True
    r_h3.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)

    table_data = [
        ["단계", "프로세스명", "입력 데이터", "수행 작업 및 데이터 변환", "출력 데이터", "수신 / 저장처"],
        ["입력", "뇌 MRI 데이터 입력", "2D-MRI 영상\n(DICOM 파일)", "의료진이 환자의 2D-MRI 영상 데이터를 시스템에 업로드", "2D-MRI 영상 데이터", "STEP 1 (전처리 모듈)"],
        ["STEP 1", "데이터 전처리\n(5단계 자동화)", "2D DICOM 영상", "1) DICOM → NIfTI 포맷 변환\n2) 두개골 제거 (Skull Stripping)\n3) 뇌 영역 공간 정규화\n4) Intensity(강도) 정규화\n5) 3D 리사이즈 (56×56×56)", "정규화된\n3D NIfTI MRI 데이터", "STEP 2 (AI 모델 모듈)"],
        ["STEP 2", "AI 모델 분석\n(심층 3D 중첩학습)", "정규화된\n3D MRI 데이터", "1) 3D-CNN: 공간 특징 추출\n2) 3D-ResNet: 심층 특징 학습\n3) CCA: 특징 결합\n4) WOA: 모델 최적화\n5) 파킨슨병 확률 산출 (정상/전구기/파킨슨병)", "- 파킨슨병 예측 확률값\n- 병변 특징 맵 데이터\n- (고위험 시) 알림 이벤트", "- STEP 3 (시각화 모듈)\n- 고위험군 알림 시스템\n- (피드백) STEP 1 재검토"],
        ["STEP 3", "진단 결과 시각화\n(M3d-CAM)", "- 예측 확률값\n- 병변 특징 맵", "M3d-CAM 기법을 사용하여 뇌 영상 내 의심 병변 영역 히트맵 시각화", "의심 병변 시각화 이미지\n(M3d-CAM 맵)", "STEP 4 (리포트 생성 모듈)"],
        ["STEP 4", "판독 리포트\n자동 생성", "- 진단 결과 요약\n- M3d-CAM 이미지\n- 분석 지표 데이터", "진단 확률 수치, 병변 시각화 맵, 분석 지표 그래프를 표준 판독 리포트로 자동 종합", "판독 리포트 데이터\n(PDF/전자 문서)", "STEP 5 (리포트 제공 모듈)"],
        ["STEP 5", "판독 리포트 제공", "판독 리포트 데이터", "의료진 다운로드용 리포트 제공 및 시스템 데이터베이스 저장", "- 최종 판독 리포트\n- DB 보관 레코드", "- 의료진 / 사용자\n- 리포트 저장 DB"]
    ]

    col_widths = [Inches(0.7), Inches(1.3), Inches(1.1), Inches(2.2), Inches(1.2), Inches(1.1)]
    
    table = doc.add_table(rows=len(table_data), cols=6)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table)

    for row_idx, row in enumerate(table.rows):
        # Header Row Styling
        is_header = (row_idx == 0)
        trPr = row._element.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>')) # Prevent row split across pages
        if is_header:
            trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))

        for col_idx, cell_value in enumerate(table_data[row_idx]):
            cell = row.cells[col_idx]
            cell.width = col_widths[col_idx]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            
            # Margins
            set_cell_margins(cell, top=100, bottom=100, left=100, right=100)
            
            # Background Color
            if is_header:
                set_cell_background(cell, "1A365D") # Deep Navy Header
            elif row_idx % 2 == 1:
                set_cell_background(cell, "FFFFFF")
            else:
                set_cell_background(cell, "F8FAFC") # Subtle zebra striping

            # Cell text format
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            
            if col_idx in [0, 1] or is_header:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT

            # Add text with proper formatting
            lines = cell_value.split('\n')
            for l_idx, line in enumerate(lines):
                if l_idx > 0:
                    p = cell.add_paragraph()
                    p.paragraph_format.space_before = Pt(0)
                    p.paragraph_format.space_after = Pt(0)
                    p.paragraph_format.line_spacing = 1.15
                    if col_idx in [0, 1] or is_header:
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    else:
                        p.alignment = WD_ALIGN_PARAGRAPH.LEFT

                run = p.add_run(line)
                run.font.name = '맑은 고딕'
                if is_header:
                    run.font.size = Pt(9.5)
                    run.font.bold = True
                    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                else:
                    run.font.size = Pt(8.5)
                    if col_idx == 0:
                        run.font.bold = True
                        run.font.color.rgb = RGBColor(0x2B, 0x6C, 0xB0)
                    elif col_idx == 1:
                        run.font.bold = True
                        run.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)
                    else:
                        run.font.color.rgb = RGBColor(0x2D, 0x37, 0x48)

    # 5. Section: 주요 처리 단계별 세부 설명
    doc.add_page_break() # Clean page break for detailed section if printing

    h4 = doc.add_paragraph()
    h4.paragraph_format.space_before = Pt(12)
    h4.paragraph_format.space_after = Pt(8)
    r_h4 = h4.add_run("4. 주요 단계별 기능 및 기술 세부사항")
    r_h4.font.name = '맑은 고딕'
    r_h4.font.size = Pt(14)
    r_h4.font.bold = True
    r_h4.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)

    steps_detail = [
        ("STEP 1. 데이터 자동 전처리 (5단계)", [
            ("DICOM → NIfTI 변환", "의료 표준 DICOM 호환성 유지 및 3D 볼륨 처리를 위한 NIfTI 파일 변환"),
            ("두개골 제거 (Skull Stripping)", "뇌실질 영역 분석의 정확도를 높이기 위해 연조직 및 두개골 영역 제거"),
            ("공간 정규화 (Spatial Normalization)", "서로 다른 환자의 뇌 크기 및 형태를 표준 템플릿 뇌 공간에 맞춰 정렬"),
            ("Intensity 정규화", "MRI 신호 강도 불균일성을 제거하고 백분위 기반 밝기 스케일 통일"),
            ("3D 리사이즈 (56×56×56)", "딥러닝 연산 효율성 확보를 위한 입체 격자 크기 재조정")
        ]),
        ("STEP 2. AI 모델 분석 (심층 3D 중첩학습)", [
            ("3D-CNN (공간 특징 추출)", "3D 뇌 MRI 영상 전체의 국소적 3D 공간 패턴 및 미세 구조 특징 추출"),
            ("3D-ResNet (심층 특징 학습)", "잔차 연결(Residual Connection)을 적용하여 그래디언트 소실 없이 깊은 층에서 고차원 파킨슨병 특징 학습"),
            ("CCA (Canonical Correlation Analysis)", "다양한 네트워크에서 추출된 뇌 특징 간 상호 연관성을 극대화하여 융합(Feature Fusion)"),
            ("WOA (Whale Optimization Algorithm)", "고래 최적화 알고리즘 기반 초매개변수 및 모델 최적화 수행"),
            ("파킨슨병 예측 및 위험군 구분", "정상(Normal), 전구기(Prodromal), 파킨슨병(Parkinson's) 3단계 확률 계산. 파킨슨병 가능성이 높은 경우 [고위험군 알림] 자동 발송")
        ]),
        ("STEP 3 ~ 5. 결과 시각화 및 판독 리포트 생성", [
            ("STEP 3. M3d-CAM 시각화", "3D Class Activation Map을 적용하여 AI가 파킨슨병으로 판단한 뇌 신경 영역의 주요 의심 병변을 히트맵 형태로 시각화"),
            ("STEP 4. 판독 리포트 자동 생성", "진단 확률 수치 (예: 정상 0.08, 전구기 0.27, 파킨슨병 0.65), 의심 병변 히트맵, 분석 지표 그래프를 자동 결합"),
            ("STEP 5. 판독 리포트 제공", "의료진용 PDF 판독 리포트 다운로드 제공 및 내역 데이터베이스 자동 저장")
        ])
    ]

    for title, items in steps_detail:
        sh = doc.add_paragraph()
        sh.paragraph_format.space_before = Pt(10)
        sh.paragraph_format.space_after = Pt(4)
        r_sh = sh.add_run(f"■ {title}")
        r_sh.font.name = '맑은 고딕'
        r_sh.font.size = Pt(11)
        r_sh.font.bold = True
        r_sh.font.color.rgb = RGBColor(0x2B, 0x6C, 0xB0)

        for sub_title, sub_desc in items:
            sp = doc.add_paragraph()
            sp.paragraph_format.space_before = Pt(0)
            sp.paragraph_format.space_after = Pt(3)
            sp.paragraph_format.left_indent = Inches(0.2)
            sp.paragraph_format.line_spacing = 1.15
            
            r_bold = sp.add_run(f"• {sub_title}: ")
            r_bold.font.name = '맑은 고딕'
            r_bold.font.size = Pt(9.5)
            r_bold.font.bold = True
            r_bold.font.color.rgb = RGBColor(0x2D, 0x37, 0x48)

            r_text = sp.add_run(sub_desc)
            r_text.font.name = '맑은 고딕'
            r_text.font.size = Pt(9.5)
            r_text.font.color.rgb = RGBColor(0x4A, 0x55, 0x68)

    # Footer section
    section = doc.sections[-1]
    footer = section.footer
    f_p = footer.paragraphs[0]
    f_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    f_run = f_p.add_run("MRI 기반 파킨슨병 조기 진단 AI 솔루션 | DFD 명세서")
    f_run.font.name = '맑은 고딕'
    f_run.font.size = Pt(8.5)
    f_run.font.color.rgb = RGBColor(0xA0, 0xAE, 0xC0)

    # Output path
    output_path = r"d:\new tensor\MRI_Parkinson_AI_Solution_DFD.docx"
    doc.save(output_path)
    print(f"Successfully generated Document at: {output_path}")

if __name__ == "__main__":
    create_document()
